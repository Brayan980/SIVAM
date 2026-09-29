"""Convierte docs/MANUAL_USUARIO.md a un archivo Word (.docx) con formato.

Uso:
    .venv\\Scripts\\python.exe scripts\\generar_manual_word.py

Genera: docs/MANUAL_USUARIO.docx
"""
import os
import re

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD_PATH = os.path.join(RAIZ, 'docs', 'MANUAL_USUARIO.md')
DOCX_PATH = os.path.join(RAIZ, 'docs', 'MANUAL_USUARIO.docx')

AZUL = RGBColor(0x1F, 0x4E, 0x79)


def agregar_runs_con_formato(parrafo, texto):
    """Interpreta **negrita** y `código` dentro de una línea."""
    # Divide por **negrita** y `codigo` conservando los delimitadores
    partes = re.split(r'(\*\*.+?\*\*|`.+?`)', texto)
    for parte in partes:
        if not parte:
            continue
        if parte.startswith('**') and parte.endswith('**'):
            run = parrafo.add_run(parte[2:-2])
            run.bold = True
        elif parte.startswith('`') and parte.endswith('`'):
            run = parrafo.add_run(parte[1:-1])
            run.font.name = 'Consolas'
            run.font.color.rgb = RGBColor(0xB0, 0x30, 0x30)
        else:
            parrafo.add_run(parte)


def limpiar_enlaces(texto):
    """Convierte [texto](enlace) en solo 'texto' y quita anclas."""
    return re.sub(r'\[(.+?)\]\(.+?\)', r'\1', texto)


def construir_documento():
    with open(MD_PATH, encoding='utf-8') as f:
        lineas = f.readlines()

    doc = Document()

    # Estilo base
    estilo = doc.styles['Normal']
    estilo.font.name = 'Calibri'
    estilo.font.size = Pt(11)

    i = 0
    en_tabla = False
    filas_tabla = []

    def volcar_tabla():
        nonlocal filas_tabla
        if not filas_tabla:
            return
        # Quita la fila separadora (---|---)
        filas = [fila for fila in filas_tabla
                 if not re.match(r'^\s*\|?[\s:\-\|]+\|?\s*$', fila)]
        celdas = [[c.strip() for c in fila.strip().strip('|').split('|')] for fila in filas]
        if not celdas:
            filas_tabla = []
            return
        ncols = max(len(f) for f in celdas)
        tabla = doc.add_table(rows=0, cols=ncols)
        tabla.style = 'Light Grid Accent 1'
        for idx, fila in enumerate(celdas):
            cells = tabla.add_row().cells
            for j in range(ncols):
                valor = fila[j] if j < len(fila) else ''
                cells[j].text = ''
                p = cells[j].paragraphs[0]
                agregar_runs_con_formato(p, limpiar_enlaces(valor))
                if idx == 0:
                    for run in p.runs:
                        run.bold = True
        doc.add_paragraph()
        filas_tabla = []

    while i < len(lineas):
        linea = lineas[i].rstrip('\n')
        strip = linea.strip()

        # Tablas
        if strip.startswith('|') and strip.endswith('|'):
            en_tabla = True
            filas_tabla.append(linea)
            i += 1
            continue
        elif en_tabla:
            volcar_tabla()
            en_tabla = False

        # Separadores horizontales
        if strip == '---':
            i += 1
            continue

        # Títulos
        if strip.startswith('#'):
            nivel = len(strip) - len(strip.lstrip('#'))
            texto = limpiar_enlaces(strip.lstrip('#').strip())
            if nivel == 1:
                p = doc.add_heading('', level=0)
                run = p.add_run(texto)
                run.font.color.rgb = AZUL
            else:
                doc.add_heading(texto, level=min(nivel, 4))
            i += 1
            continue

        # Cita / nota
        if strip.startswith('>'):
            texto = limpiar_enlaces(strip.lstrip('>').strip())
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            run_prefijo = p.add_run('  ')
            agregar_runs_con_formato(p, texto)
            for run in p.runs:
                run.italic = True
            i += 1
            continue

        # Listas numeradas
        m_num = re.match(r'^(\s*)\d+\.\s+(.*)', linea)
        if m_num:
            texto = limpiar_enlaces(m_num.group(2))
            p = doc.add_paragraph(style='List Number')
            agregar_runs_con_formato(p, texto)
            i += 1
            continue

        # Listas con viñetas
        m_bul = re.match(r'^(\s*)[-*]\s+(.*)', linea)
        if m_bul:
            indent = len(m_bul.group(1))
            texto = limpiar_enlaces(m_bul.group(2))
            estilo_lista = 'List Bullet 2' if indent >= 2 else 'List Bullet'
            try:
                p = doc.add_paragraph(style=estilo_lista)
            except KeyError:
                p = doc.add_paragraph(style='List Bullet')
            agregar_runs_con_formato(p, texto)
            i += 1
            continue

        # Línea en blanco
        if strip == '':
            i += 1
            continue

        # Párrafo normal
        p = doc.add_paragraph()
        agregar_runs_con_formato(p, limpiar_enlaces(linea))
        i += 1

    if en_tabla:
        volcar_tabla()

    doc.save(DOCX_PATH)
    print(f'Documento generado: {DOCX_PATH}')


if __name__ == '__main__':
    construir_documento()
