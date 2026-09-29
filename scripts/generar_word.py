"""Convierte un archivo Markdown de docs/ a Word (.docx) con formato.

Uso:
    .venv\\Scripts\\python.exe scripts\\generar_word.py docs\\FORMULAS_TESTS.md

Si no se pasa argumento, usa docs/FORMULAS_TESTS.md por defecto.
Genera un .docx con el mismo nombre en la misma carpeta.

Soporta: títulos (#..####), tablas, listas, citas (>), **negrita**,
`código` en línea y bloques de código con ```.
"""
import os
import re
import sys

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

AZUL = RGBColor(0x1F, 0x4E, 0x79)
GRIS_CODIGO = RGBColor(0x33, 0x33, 0x33)


def agregar_runs_con_formato(parrafo, texto):
    """Interpreta **negrita** y `código` dentro de una línea."""
    partes = re.split(r'(\*\*.+?\*\*|`.+?`)', texto)
    for parte in partes:
        if not parte:
            continue
        if parte.startswith('**') and parte.endswith('**'):
            run = parrafo.add_run(parte[2:-2])
            run.bold = True
        elif parte.startswith('`') and parte.endswith('`'):
            run = parrafo.add_run(parte[1:-1])
            run.font.name = 'Consolas'
            run.font.color.rgb = RGBColor(0xB0, 0x30, 0x30)
        else:
            parrafo.add_run(parte)


def limpiar_enlaces(texto):
    """Convierte [texto](enlace) en solo 'texto'."""
    return re.sub(r'\[(.+?)\]\(.+?\)', r'\1', texto)


def _sombrear_parrafo(parrafo, color_hex='F2F2F2'):
    """Aplica un fondo gris claro a un párrafo (para bloques de código)."""
    pPr = parrafo._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    pPr.append(shd)


def agregar_bloque_codigo(doc, lineas_codigo):
    """Agrega un bloque de código monoespaciado con fondo gris."""
    for linea in lineas_codigo:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(12)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        _sombrear_parrafo(p)
        run = p.add_run(linea if linea else ' ')
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = GRIS_CODIGO
    doc.add_paragraph()


def construir_documento(md_path):
    docx_path = os.path.splitext(md_path)[0] + '.docx'

    with open(md_path, encoding='utf-8') as f:
        lineas = f.readlines()

    doc = Document()
    estilo = doc.styles['Normal']
    estilo.font.name = 'Calibri'
    estilo.font.size = Pt(11)

    i = 0
    en_tabla = False
    filas_tabla = []

    def volcar_tabla():
        nonlocal filas_tabla
        if not filas_tabla:
            return
        filas = [fila for fila in filas_tabla
                 if not re.match(r'^\s*\|?[\s:\-\|]+\|?\s*$', fila)]
        celdas = [[c.strip() for c in fila.strip().strip('|').split('|')] for fila in filas]
        if not celdas:
            filas_tabla = []
            return
        ncols = max(len(f) for f in celdas)
        tabla = doc.add_table(rows=0, cols=ncols)
        tabla.style = 'Light Grid Accent 1'
        for idx, fila in enumerate(celdas):
            cells = tabla.add_row().cells
            for j in range(ncols):
                valor = fila[j] if j < len(fila) else ''
                cells[j].text = ''
                p = cells[j].paragraphs[0]
                agregar_runs_con_formato(p, limpiar_enlaces(valor))
                if idx == 0:
                    for run in p.runs:
                        run.bold = True
        doc.add_paragraph()
        filas_tabla = []

    while i < len(lineas):
        linea = lineas[i].rstrip('\n')
        strip = linea.strip()

        # Bloques de código con ```
        if strip.startswith('```'):
            bloque = []
            i += 1
            while i < len(lineas) and not lineas[i].strip().startswith('```'):
                bloque.append(lineas[i].rstrip('\n'))
                i += 1
            i += 1  # saltar el ``` de cierre
            if en_tabla:
                volcar_tabla()
                en_tabla = False
            agregar_bloque_codigo(doc, bloque)
            continue

        # Tablas
        if strip.startswith('|') and strip.endswith('|'):
            en_tabla = True
            filas_tabla.append(linea)
            i += 1
            continue
        elif en_tabla:
            volcar_tabla()
            en_tabla = False

        # Separadores horizontales
        if strip == '---':
            i += 1
            continue

        # Títulos
        if strip.startswith('#'):
            nivel = len(strip) - len(strip.lstrip('#'))
            texto = limpiar_enlaces(strip.lstrip('#').strip())
            if nivel == 1:
                p = doc.add_heading('', level=0)
                run = p.add_run(texto)
                run.font.color.rgb = AZUL
            else:
                doc.add_heading(texto, level=min(nivel, 4))
            i += 1
            continue

        # Cita / nota
        if strip.startswith('>'):
            texto = limpiar_enlaces(strip.lstrip('>').strip())
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            agregar_runs_con_formato(p, texto)
            for run in p.runs:
                run.italic = True
            i += 1
            continue

        # Listas numeradas
        m_num = re.match(r'^(\s*)\d+\.\s+(.*)', linea)
        if m_num:
            texto = limpiar_enlaces(m_num.group(2))
            p = doc.add_paragraph(style='List Number')
            agregar_runs_con_formato(p, texto)
            i += 1
            continue

        # Listas con viñetas
        m_bul = re.match(r'^(\s*)[-*]\s+(.*)', linea)
        if m_bul:
            indent = len(m_bul.group(1))
            texto = limpiar_enlaces(m_bul.group(2))
            estilo_lista = 'List Bullet 2' if indent >= 2 else 'List Bullet'
            try:
                p = doc.add_paragraph(style=estilo_lista)
            except KeyError:
                p = doc.add_paragraph(style='List Bullet')
            agregar_runs_con_formato(p, texto)
            i += 1
            continue

        # Línea en blanco
        if strip == '':
            i += 1
            continue

        # Párrafo normal
        p = doc.add_paragraph()
        agregar_runs_con_formato(p, limpiar_enlaces(linea))
        i += 1

    if en_tabla:
        volcar_tabla()

    doc.save(docx_path)
    print(f'Documento generado: {docx_path}')


if __name__ == '__main__':
    arg = sys.argv[1] if len(sys.argv) > 1 else os.path.join('docs', 'FORMULAS_TESTS.md')
    ruta = arg if os.path.isabs(arg) else os.path.join(RAIZ, arg)
    construir_documento(ruta)
